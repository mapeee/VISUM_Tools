# -*- coding: utf-8 -*-

import sys
import wx
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.styles import Font
from VisumPy.AddIn import AddIn, AddInState, AddInParameter
_ = AddIn.gettext


def Run(param):
    VJI = CreateVIJTable(Visum)
    lines = VJI["LINENAME"].unique().tolist()
    path = param["PuTConPath"]
    Visum.Log(20480,_("Export to: {name}").format(name = path))
    for line in lines:
        Visum.Log(20480,_("Line: {name}").format(name = line))
        VJIline = VJI[VJI["LINENAME"]==line]
        wb = Workbook()
        wb.remove(wb.active)  
        groupedVJI = VJIline.groupby("DIRECTIONCODE")
        for (direction), lineDirect_data in groupedVJI:
            ws = wb.create_sheet(title=f"{line}{direction}")
            mainline = Visum.Net.Lines.ItemByKey("3").AttValue("MAINLINENAME")
            Add2Excel([line, mainline, direction], ws, "linedir")
            lineDirect_data.sort_values(by=["DEPVJ", "VEHJOURNEYNO", "VJIINDEX"], inplace=True)
            stoparray = CreateStopList(lineDirect_data, ws)
            lineDirect_data.drop("VJIINDEX", axis=1, inplace=True)
            groupedVJ = lineDirect_data.groupby("VEHJOURNEYNO", sort=False)
            col = 3
            for (vj), vj_data in groupedVJ:
                vj_data = vj_data.reset_index()
                Add2Excel(vj, ws, "vj", col)
                mergedToStops = pd.merge(stoparray, vj_data, on='STOPPOINTNO', how='left')
                mergedToStops["EXTARRIVAL"] = mergedToStops["EXTARRIVAL"].fillna("|")
                vjstops = mergedToStops[["VJIINDEX", "EXTARRIVAL"]].values.tolist()
                Add2Excel(vjstops, ws, "vjstops", col)
                col+=1
        
        wb.save(fr"{path}\{line}.xlsx")
        del wb
    Visum.Log(20480,_("Finished"))
        
def Add2Excel(_value, _ws, method, _col=None):
    center_alignment = Alignment(horizontal='center', vertical='center')
    if method == "linedir":
        _ws.cell(row=1, column=1, value=_value[0])
        _ws.cell(row=1, column=2, value=_value[1])
        _ws.cell(row=1, column=3, value=_value[2])
        _ws.cell(row=1, column=1).font = Font(bold=True, color="e34141")
        _ws.cell(row=1, column=2).font = Font(bold=True, color="e34141")
        _ws.cell(row=1, column=3).font = Font(bold=True, color="e34141")
    if method == "stoplist":
        max_len = _value['STOP'].astype(str).str.len().max()*0.95
        _ws.cell(row=4, column=1, value="STOPNO")
        _ws.cell(row=4, column=2, value="STOP")
        _ws.column_dimensions['B'].width = max_len
        _ws.cell(row=4, column=1).font = Font(bold=True)
        _ws.cell(row=4, column=2).font = Font(bold=True)
        for i, _stop in enumerate(_value.values.tolist(), start=5):
            _ws.cell(row=i, column=1, value=_stop[1])
            _ws.cell(row=i, column=2, value=_stop[0])
            _ws.cell(row=i, column=1).font = Font(bold=True)
            _ws.cell(row=i, column=2).font = Font(bold=True)
    if method == "vj":
        _ws.cell(row=2, column=_col, value=_col-2)
        _ws.cell(row=2, column=_col).font = Font(bold=True)
        _ws.cell(row=2, column=_col).alignment = Alignment(horizontal='center', vertical='center')
        _ws.cell(row=3, column=_col, value=_value)
        _ws.cell(row=3, column=_col).font = Font(bold=True)
        _ws.cell(row=3, column=_col).alignment = Alignment(horizontal='center', vertical='center')
        _ws.cell(row=4, column=_col, value="Mo-Fr")
        _ws.cell(row=4, column=_col).font = Font(color="b2afaf")
        _ws.cell(row=4, column=_col).alignment = Alignment(horizontal='center', vertical='center')
    if method == "vjstops":
        for i in _value:
            value = i[1]
            if value != "|":
                h, m = divmod(value // 60, 60)
                value = f"{int(h):02d}:{int(m):02d}"
            _ws.cell(int(i[0]+4), int(_col), str(value))
            _ws.cell(int(i[0]+4), int(_col)).alignment = center_alignment
            if value == "|":
                _ws.cell(int(i[0]+4), int(_col)).font = Font(color="959595")
    return True

def CreateStopList(_lineDirect_data, _ws):
    _lineDirect_data.drop(["POSTLENGTH", "LINENAME", "DIRECTIONCODE"], axis=1, inplace=True)
    groupedlDd = _lineDirect_data.groupby("VEHJOURNEYNO", sort=False)
    for idx, ((VJ), VJdata) in enumerate(groupedlDd):
        if idx == 0:
            stoparray = VJdata[["STOP", "STOPPOINTNO", "VJIINDEX"]]
            continue
        mergedStopArray = pd.merge(stoparray, VJdata, on="STOPPOINTNO", how="outer")
        mergedStopArray.sort_values(by=["VJIINDEX_x", "VJIINDEX_y"], inplace=True)
        missingStops = mergedStopArray[mergedStopArray["VJIINDEX_x"].isna()].copy()
        missingStops["gap"] = missingStops["VJIINDEX_y"].diff()
        missingStops["missingGroup"] = (missingStops["gap"] > 1).cumsum() + 1
        mergedStopArray = pd.merge(mergedStopArray, missingStops[["STOPPOINTNO", "missingGroup"]], on="STOPPOINTNO", how="left")
        for i in mergedStopArray['missingGroup'].dropna().unique():
            mergedStopArray = _missingGroup(mergedStopArray, i)
        mergedStopArray["VJIINDEX"] = range(1, len(mergedStopArray) + 1)
        mergedStopArray["STOP"] = mergedStopArray["STOP_y"].fillna(mergedStopArray["STOP_x"])
        stoparray = mergedStopArray[["STOP", "STOPPOINTNO", "VJIINDEX"]]
    Add2Excel(stoparray, _ws, "stoplist")
    return stoparray

def CreateVIJTable(_Visum):
    TableFields = [["VEHJOURNEYNO", "INDEX", "EXTARRIVAL", "POSTLENGTH", r"VEHJOURNEY\LINENAME", r"VEHJOURNEY\DIRECTIONCODE",
                    r"VEHJOURNEY\DEP", r"TIMEPROFILEITEM\LINEROUTEITEM\STOPPOINT\NAME", "STOPPOINTNO"],
                   ["VEHJOURNEYNO", "VJIINDEX", "EXTARRIVAL", "POSTLENGTH", "LINENAME", "DIRECTIONCODE", "DEPVJ", "STOP", "STOPPOINTNO"]]
    _VJI = pd.DataFrame(_Visum.Net.VehicleJourneyItems.GetMultipleAttributes(TableFields[0], True),
                                      columns = TableFields[1])
    return _VJI

def _missingGroup(_df, _group):
    move_part = _df[_df['missingGroup'] == _group]
    startindex = move_part["VJIINDEX_y"].iloc[0] - 1
    rest = _df[_df['missingGroup'] != _group]
    try:
        insert_pos = rest.index.get_loc(rest[rest['VJIINDEX_y'] == startindex].index[0]) + 1
    except:
        insert_pos = 0
    df_new = pd.concat([rest.iloc[:insert_pos],move_part,rest.iloc[insert_pos:]], ignore_index=True)
    return df_new

if len(sys.argv) > 1:
    addIn = AddIn()
else:
    addIn = AddIn(Visum)

if addIn.IsInDebugMode:
    app = wx.PySimpleApp(0)
    Visum = addIn.VISUM
    addInParam = AddInParameter(addIn, None)
else:
    addInParam = AddInParameter(addIn, Parameter)

if addIn.State != AddInState.OK:
    addIn.ReportMessage(addIn.ErrorObjects[0].ErrorMessage)
else:
    try:
        defaultParam = {"Replace" : False}
        param = addInParam.Check(True, defaultParam)
        Run(param)
    except:
        addIn.HandleException(addIn.TemplateText.MainApplicationError)
